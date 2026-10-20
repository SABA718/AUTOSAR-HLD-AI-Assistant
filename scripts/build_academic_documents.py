from pathlib import Path
import re
from html import escape
from docx import Document
from docx.shared import Pt, Inches
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle, Spacer

ROOT=Path(__file__).resolve().parents[1]
DIRS=["Synopsis","Input_Data","Code","Model_Prompts_Config","Evaluation_Results","Documentation","Declarations"]
def inline(s):
    s=re.sub(r"!\[([^\]]*)\]\([^)]+\)",r"\1",s)
    s=re.sub(r"\[([^\]]+)\]\(([^)]+)\)",r"\1 (\2)",s)
    return s.replace("**","").replace("__","").replace("\x60","").strip()
def parse(s):
    lines=s.splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        i+=1
        if not line or line.startswith("\x60\x60\x60") or line in ("---","***"):continue
        if line.startswith("|"):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i+=1
            rs=[[inline(c) for c in row.strip("|").split("|")] for row in rows]
            rs=[r for r in rs if not all(re.fullmatch(r":?-{2,}:?",c or "") for c in r)]
            if rs:yield "table",rs
        elif line.startswith("#"):
            yield "head",inline(re.sub(r"^#+\s*","",line))
        elif line[:2] in ("- ","* "):
            yield "text","• "+inline(line[2:])
        else:yield "text",inline(line)
def build_word(path,items):
    d=Document()
    st=d.styles["Normal"]
    st.font.name="Times New Roman"
    st.font.size=Pt(12)
    sec=d.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.8)
    sec.left_margin=sec.right_margin=Inches(.8)
    for typ,content in items:
        if typ=="table":
            n=max(map(len,content))
            t=d.add_table(rows=0,cols=n)
            t.style="Table Grid"
            for ri,cells in enumerate(content):
                row=t.add_row()
                for j in range(n):
                    para=row.cells[j].paragraphs[0]
                    rr=para.add_run(cells[j] if j<len(cells) else "")
                    rr.font.name="Times New Roman"
                    rr.font.size=Pt(12)
                    rr.bold=(ri==0)
        else:
            p=d.add_paragraph()
            run=p.add_run(content)
            run.font.name="Times New Roman"
            run.font.size=Pt(12)
            run.bold=typ=="head"
    d.save(path.with_suffix(".docx"))
def build_pdf(path,items):
    base=ParagraphStyle("Text",fontName="TNR",fontSize=12,leading=17,spaceAfter=7,splitLongWords=1)
    hd=ParagraphStyle("Heading",parent=base,fontName="TNR-Bold",fontSize=12,spaceBefore=8)
    story=[]
    for typ,content in items:
        if typ=="table":
            n=max(map(len,content))
            rows=[[Paragraph(escape(row[j] if j<len(row) else ""),base) for j in range(n)] for row in content]
            t=Table(rows,colWidths=[(A4[0]-40*mm)/n]*n,repeatRows=1)
            t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.4,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP")]))
            story.extend([t,Spacer(1,6)])
        else:story.append(Paragraph(escape(content),hd if typ=="head" else base))
    doc=SimpleDocTemplate(str(path.with_suffix(".pdf")),pagesize=A4,leftMargin=20*mm,rightMargin=20*mm,topMargin=18*mm,bottomMargin=18*mm)
    doc.build(story)
def main():
    fontdir=Path("C:/Windows/Fonts")
    for name,file in (("TNR","times.ttf"),("TNR-Bold","timesbd.ttf")):
        font=fontdir/file
        if not font.is_file():raise RuntimeError("Actual Times New Roman font not installed: "+str(font))
        pdfmetrics.registerFont(TTFont(name,str(font)))
    count=0
    for folder in DIRS:
        for md in sorted((ROOT/folder).glob("*.md")):
            items=list(parse(md.read_text(encoding="utf-8")))
            if not items:continue
            build_word(md,items)
            build_pdf(md,items)
            count+=1
    print("Created",count,"DOCX/PDF pairs, Times New Roman 12 pt")
if __name__=="__main__":main()
